import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

BASE_DIR = r"C:\Users\Shruti\ResolveIT_Complaint_Management_System"

NAVY = RGBColor(0x1E, 0x40, 0xAF)
DARK = RGBColor(0x0F, 0x17, 0x2A)
MUTED = RGBColor(0x64, 0x74, 0x8B)

def set_cell_background(cell, hex_color):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table, color="1E40AF", sz="4", val="single"):
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders = parse_xml(f'<w:tblBorders {nsdecls("w")}><w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/><w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/><w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/><w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/><w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/><w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/></w:tblBorders>')
        tblPr[0].append(borders)

def build_viva_guide():
    doc = Document()
    
    for s in doc.sections:
        s.top_margin = Inches(1.0); s.bottom_margin = Inches(1.0)
        s.left_margin = Inches(1.0); s.right_margin = Inches(1.0)
        s.page_width = Inches(8.27); s.page_height = Inches(11.69)
        footer = s.footer
        p = footer.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r = p.add_run("ResolveIT: Viva Voce & Demonstration Guide  |  Cloud Computing Mini-Project")
        r.font.name = "Times New Roman"; r.font.size = Pt(8.5); r.font.color.rgb = MUTED

    def add_sec_title(title_text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(16); p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(title_text)
        r.font.name = "Times New Roman"; r.font.size = Pt(13); r.font.bold = True; r.font.color.rgb = NAVY

    def add_p(text, bold_prefix=None):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_before = Pt(2); p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.2
        if bold_prefix:
            rb = p.add_run(bold_prefix)
            rb.font.name = "Times New Roman"; rb.font.size = Pt(10.5); rb.font.bold = True
        r = p.add_run(text)
        r.font.name = "Times New Roman"; r.font.size = Pt(10.5)

    def add_qa(q_num, question, answer):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8); p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        r_q = p.add_run(f"Q{q_num}. {question}")
        r_q.font.name = "Times New Roman"; r_q.font.size = Pt(11); r_q.font.bold = True; r_q.font.color.rgb = NAVY
        
        p_a = doc.add_paragraph()
        p_a.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p_a.paragraph_format.space_before = Pt(1); p_a.paragraph_format.space_after = Pt(6)
        p_a.paragraph_format.line_spacing = 1.2
        r_ans_lbl = p_a.add_run("Answer: ")
        r_ans_lbl.font.name = "Times New Roman"; r_ans_lbl.font.size = Pt(10.5); r_ans_lbl.font.bold = True
        r_ans = p_a.add_run(answer)
        r_ans.font.name = "Times New Roman"; r_ans.font.size = Pt(10.5)

    # Title Header
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(10); p.paragraph_format.space_after = Pt(4)
    r = p.add_run("RESOLVEIT: VIVA VOCE & DEMONSTRATION GUIDE")
    r.font.name = "Times New Roman"; r.font.size = Pt(16); r.font.bold = True; r.font.color.rgb = NAVY

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(16)
    r = p.add_run("Course: Cloud Computing Mini-Project | Student: Shruti Pagare [Roll No: 24CO117]")
    r.font.name = "Times New Roman"; r.font.size = Pt(11); r.font.italic = True; r.font.color.rgb = MUTED

    add_sec_title("Part 1: 2-3 Minute Demonstration Script for Examiners")
    add_p(
        "Use this structured script to present your mini-project with calm authority:\n\n"
        "1. Opening (30 Seconds):\n"
        "“Good morning, respected examiners. My name is Shruti Pagare, and today I am demonstrating ResolveIT, a Cloud-Based Digital Complaint Management System. "
        "In colleges, facility breakdowns like laboratory Wi-Fi failures, projector issues, or water cooler problems are traditionally recorded in physical paper registers. "
        "These get lost, take days to process, and offer zero tracking to students. ResolveIT solves this by providing a centralized, serverless cloud portal with real-time status tracking.”\n\n"
        "2. Student Workflow Demo (60 Seconds):\n"
        "“Let me demonstrate the student experience. I log in as a student using our authentication system. On my dashboard, I can see dynamic metrics reflecting my active complaints. "
        "I click 'Lodge New Complaint', select the category 'Laboratory Equipment', enter location 'Hardware Lab', select priority 'High', and describe an oscilloscope offset error. "
        "When I submit, the system validates the form, generates a unique ID 'RES-2026-1047', and writes the record to Cloud Firestore. My dashboard immediately updates.”\n\n"
        "3. Admin Workflow & Timeline Demo (60 Seconds):\n"
        "“Now, switching to the Administrator Portal, the campus facility manager sees an oversight queue across all departments. The admin opens the newly lodged complaint, "
        "changes its status from 'Pending' to 'In Progress', and enters an official resolution note: 'Technician dispatched for calibration'. "
        "Switching back to the student view, the student opens 'Track Status' and observes our interactive visual timeline progressing to 'In Progress' with the admin's note displayed. "
        "Once work is verified, the admin marks it 'Resolved'.”\n\n"
        "4. Cloud Architecture Summary (30 Seconds):\n"
        "“From a cloud computing perspective, our frontend is decoupled and hosted on AWS Amplify / Azure Static Web Apps edge CDN, while our backend uses Firebase Authentication and Google Cloud Firestore. "
        "Granular security rules ensure students can only view their own data, and only verified administrators can update complaint statuses. Thank you, and I am ready for questions.”"
    )

    add_sec_title("Part 2: 25 Most Likely Viva Questions & Answers")

    qas = [
        ("What is Cloud Computing, and why did you use it in this project?",
         "Cloud computing is the on-demand delivery of IT resources (compute, database, storage) over the internet with pay-as-you-go pricing. We used cloud computing because it eliminates the need to buy or maintain physical servers, provides 99.99% availability, and automatically scales as complaint volumes fluctuate."),
        
        ("What are the three fundamental cloud service models? Which one does your project use?",
         "The three models are IaaS (Infrastructure as a Service), PaaS (Platform as a Service), and SaaS (Software as a Service). Our project primarily uses PaaS and BaaS (Backend as a Service)—AWS Amplify and Azure Static Web Apps provide PaaS for frontend hosting, and Firebase provides BaaS for authentication and database management."),
        
        ("What are the main cloud deployment models?",
         "Public Cloud (infrastructure owned by a provider like AWS or Google Cloud, shared across users), Private Cloud (dedicated entirely to one organization), Hybrid Cloud (combining public and private), and Community Cloud (shared among organizations with common concerns). ResolveIT runs on public cloud services."),
        
        ("Why is your frontend decoupled from the backend?",
         "Decoupling the frontend (static HTML/CSS/JS) from the backend (serverless Cloud Firestore) means the frontend can be cached globally across Content Delivery Networks (CDNs) for lightning-fast delivery (<1.2s), while database transactions execute independently via secure HTTPS API calls. This increases fault tolerance and reduces operational cost to zero."),
        
        ("What is AWS Amplify Hosting and how does it work?",
         "AWS Amplify Hosting is a fully managed cloud service for static web apps. It connects directly to a Git repository, automatically builds and deploys changes on git push, provisions free SSL certificates, and distributes the site globally using Amazon CloudFront edge locations."),
        
        ("What is Microsoft Azure Static Web Apps?",
         "Azure Static Web Apps is Microsoft's edge hosting service for modern web applications. It leverages GitHub Actions to automatically build and deploy static files to Azure's globally distributed content delivery network, providing custom domain support and built-in SSL."),
        
        ("Between AWS Amplify and Azure Static Web Apps, which do you recommend and why?",
         "We recommend Azure Static Web Apps for college mini-projects because it offers a very straightforward, beginner-friendly setup, generates an automatic GitHub Actions workflow file upon repository linking, and provides generous free allowances specifically suited for student portfolios."),
        
        ("What is Firebase Authentication and why not create custom PHP/MySQL authentication?",
         "Firebase Authentication is a managed cloud identity service. Creating custom auth often leads to security vulnerabilities like SQL injection, plain-text passwords, or broken session handling. Firebase handles password hashing, JWT token generation, session refresh, and brute-force mitigation automatically at Google's enterprise security standard."),
        
        ("What type of database is Cloud Firestore? How does it differ from SQL?",
         "Cloud Firestore is a serverless, NoSQL document-oriented database. Unlike relational databases (MySQL, PostgreSQL) which store data in rigid rows and tables, Firestore organizes data into Collections and Documents with flexible JSON-like key-value structures. It scales horizontally and supports real-time synchronization."),
        
        ("What is the difference between Authentication and Authorization in your system?",
         "Authentication confirms WHO the user is (e.g., verifying student email and password via Firebase Auth). Authorization determines WHAT the authenticated user is allowed to do (e.g., a student can only view their own complaints, whereas an administrator is authorized to update status and add resolution notes)."),
        
        ("How do you prevent ordinary students from becoming administrators in your system?",
         "In our Firebase user schema, the 'role' field is set to 'student' by default. Furthermore, our Cloud Firestore Security Rules explicitly validate that requests creating a user record cannot set role to 'admin', and update rules block students from changing their role. Administrative accounts are strictly restricted to verified emails like admin@college.edu."),
        
        ("Explain how your Cloud Firestore Security Rules work.",
         "Firestore security rules are server-side access control expressions. They inspect the request.auth token. Our rules state: allow read if request.auth.uid == resource.data.studentUid (student isolation) OR if user has admin role; allow update only if the user is a verified administrator. Unauthenticated requests are rejected outright."),
        
        ("What happens if the internet goes down during your demonstration?",
         "ResolveIT incorporates a resilient dual-mode architecture. In addition to live Firebase SDK connectivity, we have an integrated Local Cloud Simulator backed by browser localStorage with pre-seeded realistic campus complaints. This ensures that the demo works flawlessly even without internet access."),
        
        ("How does a student track the status of their complaint?",
         "Every complaint is assigned a unique ID (RES-2026-XXXX). Students open the 'Track Status' modal on their dashboard, which renders an interactive 4-step visual timeline: Submitted → Pending → In Progress → Resolved. If the administrator leaves a resolution note, it is displayed with a timestamp."),
        
        ("What are the 10 facility categories supported by ResolveIT?",
         "Hostel Facilities, Washroom & Cleanliness, Electricity, Water Supply, Classroom Facilities, Laboratory Equipment, Internet & Wi-Fi, Library, Transportation, and Other/General Inquiries."),
        
        ("What are the 4 complaint lifecycle statuses?",
         "Pending (initial state upon submission), In Progress (technician assigned or maintenance underway), Resolved (work completed and verified), and Rejected (duplicate or invalid inquiry with logged reason)."),
        
        ("How does the system calculate dashboard statistics?",
         "Statistics are computed dynamically from actual database records, not hardcoded. The application queries the complaints collection, counts the array length for Total, and filters by status attribute to display live counts for Pending, In Progress, and Resolved."),
        
        ("What is a Content Delivery Network (CDN) and why is it important in cloud hosting?",
         "A CDN is a geographically distributed network of proxy servers that cache web assets close to end users. Both AWS Amplify (CloudFront) and Azure SWA use CDNs, ensuring that students and staff load the webpage in sub-second speeds regardless of geographic location."),
        
        ("What is High Availability in cloud computing?",
         "High availability means the system remains operational and accessible for a high percentage of time (e.g., 99.99% uptime) through automated redundancy and failover mechanisms provided by cloud providers across multiple Availability Zones."),
        
        ("What is Scalability? How does your project achieve it?",
         "Scalability is the ability of a system to handle growing workloads by adding resources. Because our frontend is static on a CDN and Cloud Firestore scales automatically without server bottlenecks, ResolveIT can handle 10 or 10,000 complaints simultaneously without crashing."),
        
        ("What is Serverless computing?",
         "Serverless is a cloud execution model where the cloud provider dynamically manages the allocation of servers. Developers do not provision, scale, or maintain virtual machines; they only write application code and pay for actual execution resources consumed."),
        
        ("Can a student edit or delete a complaint after submitting it?",
         "By design and security policy, students cannot delete or alter a complaint once submitted to maintain an unalterable audit trail. Only the administrator can update the status or append resolution remarks."),
        
        ("What are the primary limitations of ResolveIT in its current version?",
         "Currently, image attachments are not uploaded directly to Firebase Storage (only text descriptions), and automated email/SMS notifications are not yet implemented (would require Firebase Cloud Functions)."),
        
        ("What are the future enhancements planned for this project?",
         "Adding photo attachments via Firebase Cloud Storage, automated push/email notifications via Cloud Functions, a cross-platform Flutter mobile app, and AI-driven automated routing of complaints directly to technician phone numbers."),
        
        ("Why is this mini-project appropriate for a Cloud Computing course?",
         "Because it directly demonstrates practical cloud concepts: PaaS/BaaS integration, decoupled edge hosting, NoSQL cloud database design, tokenized cloud authentication, cloud security rules, high availability, and automated CI/CD deployment pipelines.")
    ]

    for idx, (q, a) in enumerate(qas):
        add_qa(idx + 1, q, a)

    add_sec_title("Part 3: Pre-Viva Testing Checklist")
    add_p(
        "Complete this 5-minute checklist before presenting to the examiner:\n"
        "• [ ] Open index.html in Google Chrome or Microsoft Edge — ensure styles load cleanly.\n"
        "• [ ] Test Student Quick Login ('👤 Student Demo') — verifies redirect to dashboard.html.\n"
        "• [ ] Test 'Lodge New Complaint' — submit a test issue and verify that the Total counter increases by 1.\n"
        "• [ ] Click 'Track Status' on the newly submitted issue — verify the timeline renders 'Pending'.\n"
        "• [ ] Test Admin Quick Login ('🛡️ Admin Demo') — verifies redirect to admin.html.\n"
        "• [ ] In Admin Portal, find the test issue, change status to 'In Progress', enter a note, and save.\n"
        "• [ ] Return to Student view — verify that the timeline updates to 'In Progress' and displays the note.\n"
        "• [ ] Test search input and category dropdown — confirm real-time filtering.\n"
        "• [ ] Test Logout button — verifies clean return to login.html."
    )

    add_sec_title("Part 4: Common Errors and Troubleshooting")
    add_p(
        "1. Issue: Pages appear unstyled or look like raw HTML.\n"
        "   Cause: Incorrect relative path to css/style.css.\n"
        "   Fix: Ensure the browser opens index.html from within the ResolveIT_Complaint_Management_System directory.\n\n"
        "2. Issue: Clicking Submit Complaint does not update the table.\n"
        "   Cause: Browser JavaScript is disabled or form validation failed on a blank required field.\n"
        "   Fix: Ensure Title, Category, Location, and Description are filled. Check browser Developer Console (F12) for messages.\n\n"
        "3. Issue: Firebase console shows 'Permission Denied'.\n"
        "   Cause: Firestore security rules are blocking the read/write request because user is unauthenticated.\n"
        "   Fix: Ensure user is logged in, or paste the provided firestore.rules into your Firebase Console Rules tab."
    )

    out_file = os.path.join(BASE_DIR, "ResolveIT_Viva_and_Demo_Guide.docx")
    doc.save(out_file)
    print("Viva guide generated successfully:", out_file)

if __name__ == "__main__":
    build_viva_guide()
